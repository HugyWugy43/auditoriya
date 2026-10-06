#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Скрипт для конвертации отчёта из Markdown в DOCX формат
Требуется: pip install python-docx
"""

import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def parse_markdown_to_docx(md_file, docx_file):
    """Конвертирует Markdown файл в DOCX"""
    
    doc = Document()
    
    # Настройка стилей по умолчанию
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    
    # Читаем Markdown файл
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    lines = content.split('\n')
    i = 0
    
    while i < len(lines):
        line = lines[i].strip()
        
        # Пропускаем пустые строки и разделители
        if not line or line == '---':
            i += 1
            continue
        
        # Заголовок уровня 1 (##)
        if line.startswith('## ') and not line.startswith('###'):
            text = line[3:].strip()
            p = doc.add_heading(text, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            # Устанавливаем шрифт Times New Roman для заголовков
            for run in p.runs:
                run.font.name = 'Times New Roman'
            i += 1
            continue
        
        # Заголовок уровня 2 (###)
        if line.startswith('### ') and not line.startswith('####'):
            text = line[4:].strip()
            p = doc.add_heading(text, level=2)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Times New Roman'
            i += 1
            continue
        
        # Заголовок уровня 3 (####)
        if line.startswith('#### '):
            text = line[5:].strip()
            p = doc.add_heading(text, level=3)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Times New Roman'
            i += 1
            continue
        
        # Код блоки (```)
        if line.startswith('```'):
            # Пропускаем строку с ```
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            # Пропускаем закрывающий ```
            if i < len(lines):
                i += 1
            
            # Добавляем код как параграф с моноширинным шрифтом
            code_text = '\n'.join(code_lines).strip()
            if code_text:
                p = doc.add_paragraph(code_text)
                p.style = 'No Spacing'
                for run in p.runs:
                    run.font.name = 'Courier New'
                    run.font.size = Pt(11)
            continue
        
        # Списки с маркерами
        if line.startswith('- ') or line.startswith('* '):
            text = line[2:].strip()
            # Удаляем жирный текст маркеры (**)
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            # Удаляем обратные кавычки
            text = re.sub(r'`([^`]+)`', r'\1', text)
            p = doc.add_paragraph(text, style='List Bullet')
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(14)
            i += 1
            continue
        
        # Нумерованные списки
        if re.match(r'^\d+\.\s', line):
            text = re.sub(r'^\d+\.\s', '', line)
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'`([^`]+)`', r'\1', text)
            p = doc.add_paragraph(text, style='List Number')
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(14)
            i += 1
            continue
        
        # Обычный текст
        text = line
        # Удаляем обратные кавычки (inline code)
        text = re.sub(r'`([^`]+)`', r'\1', text)
        # Удаляем ссылки [text](url) -> text
        text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)
        
        if text.strip():
            p = doc.add_paragraph()
            # Обрабатываем жирный текст (**text**)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(14)
                elif part.strip():
                    run = p.add_run(part)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(14)
        
        i += 1
    
    # Сохраняем документ
    doc.save(docx_file)
    print(f"✓ Отчёт успешно сохранён в {docx_file}")

if __name__ == '__main__':
    try:
        parse_markdown_to_docx('REPORT.md', 'REPORT.docx')
        print("\nФайл REPORT.docx создан и готов к использованию!")
    except ImportError:
        print("Ошибка: Не установлена библиотека python-docx")
        print("Установите её командой: pip install python-docx")
    except FileNotFoundError as e:
        print(f"Ошибка: Файл не найден - {e}")
    except Exception as e:
        print(f"Ошибка при конвертации: {e}")
        import traceback
        traceback.print_exc()
